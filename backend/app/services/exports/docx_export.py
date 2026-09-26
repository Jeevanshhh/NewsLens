"""Word (.docx) export using python-docx.

Landscape document with a title, search parameters, category/state/provider
summary tables and a full article table. All counts are derived from the rows
actually passed in.
"""
from __future__ import annotations

import io
from typing import List

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt

from app.services.exports.common import (
    EXPORT_COLUMNS,
    HEADERS,
    cell,
    search_parameter_rows,
    summarize,
)


def _set_landscape(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    new_width, new_height = section.page_height, section.page_width
    section.page_width = new_width
    section.page_height = new_height


def _fill_table(doc: Document, headers: List[str], rows_data: List[List[str]]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = text
        for paragraph in hdr[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for row in rows_data:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = "" if value is None else str(value)
    return table


def build_docx(rows: List[dict], meta: dict) -> bytes:
    doc = Document()
    _set_landscape(doc)

    doc.add_heading("NewsLens - Search Export", level=0)

    doc.add_heading("Search Parameters", level=1)
    _fill_table(doc, ["Parameter", "Value"], [list(p) for p in search_parameter_rows(meta)])

    doc.add_heading("Summaries", level=1)
    for title, key in (("By Category", "category"), ("By State", "state"), ("By Provider", "provider")):
        doc.add_heading(title, level=2)
        pairs = summarize(rows, key)
        if pairs:
            _fill_table(doc, [title.replace("By ", ""), "Count"], [[k, str(v)] for k, v in pairs])
        else:
            doc.add_paragraph("No data.")

    doc.add_heading(f"Articles ({len(rows)})", level=1)
    if rows:
        body = [[cell(r, key) for _, key in EXPORT_COLUMNS] for r in rows]
        table = _fill_table(doc, HEADERS, body)
        for row in table.rows:
            for c in row.cells:
                for paragraph in c.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(7)
    else:
        doc.add_paragraph("No articles matched this search.")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
