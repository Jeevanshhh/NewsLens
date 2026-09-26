"""Excel (.xlsx) export with multiple summary sheets."""
from __future__ import annotations

import io
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Font

from app.services.exports.common import HEADERS, KEYS, search_parameter_rows, summarize


def _style_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True)


def _add_table(ws, headers, data_rows):
    ws.append(headers)
    for row in data_rows:
        ws.append(row)
    _style_header(ws)


def build_xlsx(rows: List[dict], meta: dict) -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = "Articles"
    _add_table(ws, HEADERS, [[r.get(k, "") for k in KEYS] for r in rows])

    for title, key in (
        ("Category Summary", "category"),
        ("State Summary", "state"),
        ("Source Summary", "provider"),
    ):
        s = wb.create_sheet(title)
        _add_table(s, ["Value", "Count"], summarize(rows, key))

    params = wb.create_sheet("Search Parameters")
    _add_table(params, ["Parameter", "Value"], search_parameter_rows(meta))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
